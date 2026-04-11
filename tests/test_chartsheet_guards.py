import json

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from excel_mcp.calculations import apply_formula
from excel_mcp.chart import create_chart_in_sheet
from excel_mcp.data import append_table_rows, update_rows_by_key, write_data
from excel_mcp.exceptions import (
    DataError,
    SheetError,
    ValidationError,
)
from excel_mcp.formatting import format_range
from excel_mcp.server import (
    create_table as create_table_tool,
    get_data_validation_info as get_data_validation_info_tool,
)
from excel_mcp.sheet import (
    copy_range_operation,
    copy_sheet,
    get_sheet_protection,
    set_column_widths,
    set_freeze_panes,
    set_print_area,
)
from excel_mcp.tables import create_excel_table
from excel_mcp.validation import (
    validate_formula_in_cell_operation,
    validate_range_in_sheet_operation,
)


def _create_chart_workbook(tmp_path, *, chartsheet_first: bool = False) -> str:
    filepath = str(tmp_path / "chartsheet-guards.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [("Name", "Value"), ("Alice", 30), ("Bob", 25)]:
        ws.append(row)

    chart = BarChart()
    data = Reference(ws, min_col=2, min_row=1, max_row=3)
    categories = Reference(ws, min_col=1, min_row=2, max_row=3)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    chart_sheet = wb.create_chartsheet("Charts")
    chart_sheet.add_chart(chart)

    if chartsheet_first:
        wb._sheets = [chart_sheet, ws]
        wb.active = 0

    wb.save(filepath)
    wb.close()
    return filepath


def test_validation_operations_reject_chart_sheets_with_clear_errors(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        validate_formula_in_cell_operation(filepath, "Charts", "A1", "=1+1")

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        validate_range_in_sheet_operation(filepath, "Charts", "A1", "B2")


def test_apply_formula_rejects_chart_sheet_with_clear_error(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        apply_formula(filepath, "Charts", "A1", "=1+1")


def test_format_range_rejects_chart_sheet_with_clear_error(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        format_range(filepath, "Charts", "A1", "B2", bold=True)


def test_sheet_operations_reject_chart_sheets_with_clear_errors(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    with pytest.raises(SheetError, match="Sheet 'Charts' is a chartsheet"):
        copy_sheet(filepath, "Charts", "Charts Copy")

    with pytest.raises(SheetError, match="Sheet 'Charts' is a chartsheet"):
        get_sheet_protection(filepath, "Charts")

    with pytest.raises(SheetError, match="Sheet 'Charts' is a chartsheet"):
        set_print_area(filepath, "Charts", "A1:B2")

    with pytest.raises(SheetError, match="Sheet 'Charts' is a chartsheet"):
        set_column_widths(filepath, "Charts", {"A": 14})

    with pytest.raises(SheetError, match="Sheet 'Charts' is a chartsheet"):
        set_freeze_panes(filepath, "Charts", "B2")

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        copy_range_operation(filepath, "Charts", "A1", "B2", "D1")


def test_create_chart_rejects_chart_sheet_targets_and_sources(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        create_chart_in_sheet(filepath, "Charts", "Data!A1:B3", "bar", "D1")

    with pytest.raises(ValidationError, match="Sheet 'Charts' is a chartsheet"):
        create_chart_in_sheet(filepath, "Data", "Charts!A1:B3", "bar", "D1")


def test_write_operations_reject_chart_sheets_with_clear_errors(tmp_path):
    filepath = _create_chart_workbook(tmp_path, chartsheet_first=True)

    with pytest.raises(DataError, match="Sheet 'Charts' is a chartsheet"):
        write_data(filepath, "Charts", [["Mallory"]], "A1", dry_run=True)

    with pytest.raises(DataError, match="Active sheet 'Charts' is a chartsheet"):
        write_data(filepath, None, [["Mallory"]], "A1", dry_run=True)

    with pytest.raises(DataError, match="Sheet 'Charts' is a chartsheet"):
        append_table_rows(
            filepath,
            "Charts",
            [{"Name": "Mallory", "Value": 44}],
            dry_run=True,
        )

    with pytest.raises(DataError, match="Sheet 'Charts' is a chartsheet"):
        update_rows_by_key(
            filepath,
            "Charts",
            "Name",
            [{"Name": "Alice", "Value": 31}],
            dry_run=True,
        )


def test_table_creation_rejects_chart_sheets_with_clear_errors(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    with pytest.raises(DataError, match="Sheet 'Charts' is a chartsheet"):
        create_excel_table(filepath, "Charts", "A1:B2")

    payload = json.loads(create_table_tool(filepath, "Charts", "A1:B2"))
    assert payload["ok"] is False
    assert payload["operation"] == "create_table"
    assert "Sheet 'Charts' is a chartsheet" in payload["error"]["message"]


def test_validation_info_tool_rejects_chart_sheets_with_clear_error(tmp_path):
    filepath = _create_chart_workbook(tmp_path)

    payload = json.loads(get_data_validation_info_tool(filepath, "Charts"))
    assert payload["ok"] is False
    assert payload["operation"] == "get_data_validation_info"
    assert payload["error"]["type"] == "SheetError"
    assert "Sheet 'Charts' is a chartsheet" in payload["error"]["message"]
